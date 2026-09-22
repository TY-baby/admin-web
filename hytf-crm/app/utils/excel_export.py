from io import BytesIO
from typing import Iterable, Sequence
from openpyxl import Workbook


def export_rows_to_xlsx(headers: Sequence[str], rows: Iterable[Sequence]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_rows_to_xlsx_merged(headers: Sequence[str], rows: Iterable[Sequence],
                               merge_groups: Sequence[tuple] = (),
                               merge_cols: Sequence[int] = ()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    for s, e in merge_groups:
        if e > s:
            for c in merge_cols:
                ws.merge_cells(start_row=s + 2, start_column=c + 1,
                               end_row=e + 2, end_column=c + 1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
